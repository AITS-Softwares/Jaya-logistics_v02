from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "output/Jaya_Logistics_Driver_App_User_Guide.docx"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), fill)
    tc_pr.append(shade)

def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag); borders.append(node)
        node.set(qn('w:val'), 'single'); node.set(qn('w:sz'), '6'); node.set(qn('w:color'), color)

def set_font(run, size=11, bold=False, color=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def add_text(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); set_font(r, 11, True)
        r = p.add_run(text[len(bold_lead):]); set_font(r)
    else:
        r = p.add_run(text); set_font(r)
    p.paragraph_format.space_after = Pt(7)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text); set_font(r)
    p.paragraph_format.space_after = Pt(4)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8); section.right_margin = Inches(0.8)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(11)
for name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 12)]:
    styles[name].font.name = 'Aptos'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run('Jaya Logistics Driver App User Guide'); set_font(r, 22, True)
subtitle = doc.add_paragraph(); r = subtitle.add_run('For drivers and dispatch team'); set_font(r, 11, False, '555555')
subtitle.paragraph_format.space_after = Pt(16)

add_text(doc, 'This guide explains how a driver uses the Jaya Logistics Driver app for assigned delivery trips. The app shows only the driver’s own trips and shares location only while an active trip is being tracked.')

h = doc.add_paragraph(style='Heading 1'); r = h.add_run('Before a Driver Can Use the App'); set_font(r, 15, True)
add_bullet(doc, 'The company must install the APK on the Android phone.')
add_bullet(doc, 'Dispatch must create the driver profile, set the driver mobile number and PIN, and assign the driver to a delivery trip in the ERP.')
add_bullet(doc, 'The phone must have mobile data or Wi-Fi, GPS enabled, and a working HTTPS connection to the company server.')
add_bullet(doc, 'The driver must not share the PIN with anyone.')

h = doc.add_paragraph(style='Heading 1'); r = h.add_run('Driver Steps'); set_font(r, 15, True)
steps = [
    ('1. Open the app', 'Tap Jaya Logistics Driver after installation.'),
    ('2. Sign in', 'Enter the assigned mobile number and PIN, then tap Sign in securely.'),
    ('3. Wait for approval', 'On first sign-in the phone is registered. The driver must wait while dispatch approves this device in the ERP, then tap Refresh approval.'),
    ('4. Check assigned trip', 'After approval, only active trips assigned to that driver appear. The driver cannot see other ERP screens or other drivers’ trips.'),
    ('5. Start the trip', 'Tap the available trip action such as Start trip, At pickup, Start delivery, or Arrived at stop as the delivery progresses.'),
    ('6. Allow permissions', 'When location tracking starts, allow Precise Location and Notifications. In Android Settings for the app, set Location to Allow all the time so tracking can continue during delivery.'),
    ('7. Complete the trip', 'Use the available delivery and completion actions only after the delivery is finished. Tracking stops when the trip is completed.')
]
table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
table.rows[0].cells[0].text = 'Step'; table.rows[0].cells[1].text = 'What to do'
for cell in table.rows[0].cells:
    set_cell_shading(cell, '1F2937'); set_cell_border(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs: set_font(run, 10, True, 'FFFFFF')
for i, (step, detail) in enumerate(steps):
    cells = table.add_row().cells
    for cell, text in zip(cells, (step, detail)):
        cell.text = text; set_cell_border(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i % 2 == 1: set_cell_shading(cell, 'F5F7FA')
        for run in cell.paragraphs[0].runs: set_font(run, 10, cell is cells[0])

h = doc.add_paragraph(style='Heading 1'); r = h.add_run('Permissions and Phone Settings'); set_font(r, 15, True)
add_text(doc, 'The app needs only the permissions required for delivery tracking. It does not require access to the driver’s camera, contacts, or files for the current tracking workflow.')
perm = doc.add_table(rows=1, cols=2); perm.alignment = WD_TABLE_ALIGNMENT.CENTER; perm.style = 'Table Grid'
perm.rows[0].cells[0].text = 'Permission or setting'; perm.rows[0].cells[1].text = 'Why it is needed'
for cell in perm.rows[0].cells:
    set_cell_shading(cell, '1F2937'); set_cell_border(cell)
    for run in cell.paragraphs[0].runs: set_font(run, 10, True, 'FFFFFF')
for i, row in enumerate([
    ('Precise location', 'Records delivery movement and arrival or departure events.'),
    ('Allow all the time location', 'Allows location updates while the app is not open during an active trip.'),
    ('Notifications', 'Shows that delivery tracking is active in the notification bar.'),
    ('Mobile data or Wi-Fi', 'Sends trip actions and location updates to the server.'),
    ('GPS / Location services', 'Must remain switched on while delivery tracking is required.'),
    ('Battery setting', 'Set the app to Unrestricted or do not optimize it, if the phone stops tracking in the background.')
]):
    cells = perm.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text; set_cell_border(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i % 2 == 1: set_cell_shading(cell, 'F5F7FA')
        for run in cell.paragraphs[0].runs: set_font(run, 10, cell is cells[0])

h = doc.add_paragraph(style='Heading 1'); r = h.add_run('Dispatch Team Steps'); set_font(r, 15, True)
add_bullet(doc, 'Create or confirm the driver profile and PIN in the ERP.')
add_bullet(doc, 'Assign the driver to the relevant Loading Info or delivery trip before asking the driver to sign in.')
add_bullet(doc, 'After first sign-in, open the Driver Mobile area in the ERP and approve the registered device.')
add_bullet(doc, 'Monitor active trips and live location through the ERP tracking screens.')
add_bullet(doc, 'If a phone is lost, replaced, or no longer authorized, revoke its device approval in the ERP.')

h = doc.add_paragraph(style='Heading 1'); r = h.add_run('Quick Troubleshooting'); set_font(r, 15, True)
troubles = [
    ('No trips displayed', 'Ask dispatch to confirm the device is approved and an active trip is assigned to the correct driver.'),
    ('Location not updating', 'Turn on GPS, check mobile data, allow Precise Location and Allow all the time, then remove battery optimization for the app.'),
    ('Cannot sign in', 'Check the mobile number and PIN with dispatch, then confirm the server URL is reachable over HTTPS.'),
    ('New or replaced phone', 'Sign in on the new phone and ask dispatch to approve that device. The old device should be revoked.')
]
for title, body in troubles:
    add_text(doc, title + ': ' + body, bold_lead=title + ':')

h = doc.add_paragraph(style='Heading 1'); r = h.add_run('Important Operating Rules'); set_font(r, 15, True)
add_bullet(doc, 'Use the app only for the trip assigned to the signed-in driver.')
add_bullet(doc, 'Keep GPS, data connection, and battery available during active delivery tracking.')
add_bullet(doc, 'Report incorrect trip details or a tracking problem to dispatch immediately.')
add_bullet(doc, 'The app is for driver operations only; full ERP modules remain accessible only through authorized ERP users.')

footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Jaya Logistics Driver App'); set_font(r, 9, False, '666666')

doc.save(OUT)
